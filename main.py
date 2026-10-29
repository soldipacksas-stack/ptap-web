from fastapi import FastAPI, Depends, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, FileResponse
from fastapi.templating import Jinja2Templates
from pydantic import BaseModel
from sqlalchemy.orm import Session
from datetime import datetime
from typing import List, Optional
import openpyxl
import os

from database import engine, Base, get_db
import models

Base.metadata.create_all(bind=engine)

app = FastAPI(title="Control Operacional PTAP - Web")
templates = Jinja2Templates(directory="templates")

# Lista oficial de operadores PTAP
OPERADORES_PTAP = [
    "BALLEN MERCHAN DIEGO ARMANDO",
    "CORDOBA JIMENEZ HECTOR ENRIQUE",
    "GARCIA SANCHEZ RICARDO ANDRES",
    "GOMEZ CACERES RICARDO ALFONSO",
    "MELO DIAZ BRYAM ANDRES",
    "MENDOZA CHARRY KEVIN JHOJAN",
    "OSPINA HERNANDEZ JORGE ALEJANDRO",
    "ROBAYO CASTRO DANIELA",
    "RODRIGUEZ PUENTES JAIME ANDRES",
    "RODRIGUEZ VELANDIA JENNY PAOLA",
    "SOLIS EPES JAVIER MAURICIO",
    "ARTENGA SANDOVAL NELSON DAVID",
    "PAMPLONA FUQUENE JULIAN DANIEL"
]

# Maestra del Plan de Calidad PTAP (99 Registros Oficiales validados con Excel)
MAESTRA_PLAN_CALIDAD = [
    ("Entrada Agua Potable", "Agua Fuente (Río)", "Todos", "Todos", "%"),
    ("Entrada Agua Potable", "Agua Fuente (Pozo 1)", "Todos", "Todos", "%"),
    ("Entrada Agua Potable", "Agua Fuente (Pozo 2)", "Todos", "Todos", "%"),
    ("Entrada Agua Potable", "Agua Fuente (Pozo 3)", "Todos", "Todos", "%"),
    ("Entrada Agua Potable", "Cloro residual (Muestra 1)", "Todos", "Todos", "0,5 - 2"),
    ("Entrada Agua Potable", "Cloro residual (Muestra 2)", "Todos", "Todos", "0,5 - 2"),
    ("Entrada Agua Potable", "pH", "Todos", "Todos", "6,5 - 7,5"),
    ("Entrada Agua Potable", "Hierro total (Fe)", "Todos", "Turno 1", "< 0,10"),
    ("Entrada Agua Potable", "Evaluación Sensorial", "Todos", "Todos", "Pasa / No pasa"),
    ("Entrada Agua Potable", "Turbiedad (Lab)", "Todos", "Turno 1", "< 0,50 NTU"),
    ("Entrada Agua Potable", "Turbiedad (Linea)", "Todos", "Turno 1", "< 0,50 NTU"),
    ("Entrada Agua Potable", "Sólidos disueltos totales (TDS)", "Miércoles", "Turno 1", "< 500 ppm"),
    ("Entrada Agua Potable", "Conductividad", "Miércoles", "Turno 1", "< 745 us/cm"),
    ("Entrada Agua Potable", "Color", "Miércoles", "Turno 1", "< 15 UPC"),
    ("Entrada Agua Potable", "Dureza Total", "Miércoles", "Turno 1", "< 100 ppm"),
    ("Entrada Agua Potable", "Cloruros", "Miércoles", "Turno 1", "< 250 ppm"),
    ("Entrada Agua Potable", "Diferencial de Presión Filtro 1", "Miércoles", "Turno 1", "< 10 PSI"),
    ("Entrada Agua Potable", "Diferencial de Presión Filtro 2", "Miércoles", "Turno 1", "< 10 PSI"),
    ("Entrada Agua Potable", "Sulfatos", "Martes", "Turno 1", "< 50 ppm"),
    ("Entrada Agua Potable", "Aluminio (Al)", "Martes", "Turno 1", "< 0,10 ppm"),
    ("Entrada Agua Potable", "Alcalinidad Total", "Martes", "Turno 1", "< 90 ppm"),
    ("Entrada Agua Potable", "Manganeso", "Martes", "Turno 1", "< 0,050 ppm"),
    ("Agua potable para Servicios Generales", "Cloro residual", "Todos", "Turno 1", "0,5 - 2"),
    ("Casino principal", "Cloro residual", "Jueves", "Turno 1", "0,5 - 2"),
    ("Baño envasado", "Cloro residual", "Jueves", "Turno 1", "0,5 - 2"),
    ("Casino Postobón", "Cloro residual", "Jueves", "Turno 1", "0,5 - 2"),
    ("Agua Suavizada", "Dureza Total", "Todos", "Turno 1", "0 - 5"),
    ("Agua Suavizada", "pH", "Todos", "Turno 1", "6,5 - 8"),
    ("Agua Suavizada", "Conductividad", "Miércoles", "Turno 1", "< 745 us/cm"),
    ("Agua Suavizada", "Hierro total (Fe)", "Miércoles", "Turno 1", "< 0,10"),
    ("Agua Suavizada", "Color", "Miércoles", "Turno 1", "< 15 UPC"),
    ("Agua Suavizada", "Solidos disueltos totales (TDS)", "Miércoles", "Turno 1", "< 500 ppm"),
    ("Almacenamiento de agua potable para proceso", "Cloro residual", "Todos", "Todos", "0,5 - 2"),
    ("Ultrafiltración 1 agua Cervezera", "Turbiedad", "Todos", "Turno 1", "< 0,50 NTU"),
    ("Ultrafiltración 1 agua Cervezera", "Hierro total (Fe)", "Miércoles", "Turno 1", "< 0,10"),
    ("Ultrafiltración 1 agua Cervezera", "Color", "Miércoles", "Turno 1", "< 15 UPC"),
    ("Ultrafiltración 1 agua Cervezera", "Manganeso", "Martes", "Turno 1", "< 0,050 ppm"),
    ("Ultrafiltración 2 agua Cervezera", "Turbiedad", "Todos", "Turno 1", "< 0,50 NTU"),
    ("Ultrafiltración 2 agua Cervezera", "Hierro total (Fe)", "Miércoles", "Turno 1", "< 0,10"),
    ("Ultrafiltración 2 agua Cervezera", "Color", "Miércoles", "Turno 1", "< 15 UPC"),
    ("Ultrafiltración 2 agua Cervezera", "Manganeso", "Martes", "Turno 1", "< 0,050 ppm"),
    ("Ultrafiltración 1 de agua de Servicios", "Turbiedad", "Todos", "Turno 1", "< 0,50 NTU"),
    ("Ultrafiltración 1 de agua de Servicios", "Hierro total (Fe)", "Miércoles", "Turno 1", "< 0,10"),
    ("Ultrafiltración 1 de agua de Servicios", "Color", "Miércoles", "Turno 1", "< 15 UPC"),
    ("Ultrafiltración 1 de agua de Servicios", "Manganeso", "Martes", "Turno 1", "< 0,050 ppm"),
    ("Ultrafiltración 2 de agua de Servicios", "Turbiedad", "Todos", "Turno 1", "< 0,50 NTU"),
    ("Ultrafiltración 2 de agua de Servicios", "Hierro total (Fe)", "Miércoles", "Turno 1", "< 0,10"),
    ("Ultrafiltración 2 de agua de Servicios", "Color", "Miércoles", "Turno 1", "< 15 UPC"),
    ("Ultrafiltración 2 de agua de Servicios", "Manganeso", "Martes", "Turno 1", "< 0,050 ppm"),
    ("Tanque de almacenamiento PTAP cervecera", "Cloro residual", "Todos", "Todos", "0,5 - 2"),
    ("Tanque de almacenamiento PTAP Servicios", "Cloro residual", "Todos", "Todos", "0,5 - 2"),
    ("Filtracion por carbon agua cervecera", "Cloro residual F C1", "Martes", "Turno 1", "< 0,10"),
    ("Filtracion por carbon agua cervecera", "Cloro residual F C2", "Martes", "Turno 1", "< 0,10"),
    ("Filtracion por carbon agua cervecera", "Cloro residual F C3", "Martes", "Turno 1", "< 0,10"),
    ("Filtracion por carbon agua de servicios", "Cloro residual F C1", "Martes", "Turno 1", "< 0,10"),
    ("Filtracion por carbon agua de servicios", "Cloro residual F C2", "Martes", "Turno 1", "< 0,10"),
    ("Filtracion por carbon agua de servicios", "Cloro residual F C3", "Martes", "Turno 1", "< 0,10"),
    ("Salida tren filtros agua cervecera", "Cloro residual (Lab)", "Todos", "Turno 1", "< 0,10"),
    ("Salida tren filtros agua cervecera", "Cloro residual (Linea)", "Todos", "Turno 1", "< 0,10"),
    ("Salida tren filtros agua cervecera", "Concentración dioxido de cloro (Lab)", "Todos", "Turno 1", "0,4 - 0,8"),
    ("Salida tren filtros agua cervecera", "Concentración dioxido de cloro (Linea)", "Todos", "Turno 1", "0,4 - 0,8"),
    ("Salida tren filtros agua de servicios", "Cloro residual (Lab)", "Todos", "Turno 1", "< 0,10"),
    ("Salida tren filtros agua de servicios", "Cloro residual (Linea)", "Todos", "Turno 1", "< 0,10"),
    ("Salida tren filtros agua de servicios", "Concentración dioxido de cloro (Lab)", "Todos", "Turno 1", "0,4 - 0,8"),
    ("Salida tren filtros agua de servicios", "Concentración dioxido de cloro (Linea)", "Todos", "Turno 1", "0,4 - 0,8"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "pH (Muestra 1)", "Todos", "Todos", "6,5 - 7,2"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "pH (Muestra 2)", "Todos", "Todos", "6,5 - 7,2"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Cloro residual", "Todos", "Turno 1", "< 0,10"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Turbiedad", "Todos", "Turno 1", "< 0,50"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Evaluación Sensorial", "Todos", "Todos", "Pasa / No Pasa"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Intensidad lampara UV 1", "Todos", "Todos", "> 25"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Intensidad lampara UV 2", "Todos", "Todos", "> 25"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Color", "Miércoles", "Turno 1", "< 15 UPC"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Hierro total (Fe)", "Miércoles", "Turno 1", "< 0,10"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Dureza Total", "Miércoles", "Turno 1", "< 100 ppm"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Cloruros", "Miércoles", "Turno 1", "< 70 ppm"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Solidos disueltos totales", "Miércoles", "Turno 1", "< 500 ppm"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Conductividad", "Miércoles", "Turno 1", "< 745 us/cm"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Sulfatos", "Martes", "Turno 1", "< 50 ppm"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Manganeso", "Martes", "Turno 1", "< 0,050 ppm"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Alcalinidad Total", "Martes", "Turno 1", "< 90 ppm"),
    ("Tratamiento con luz Ultravioleta agua cervecera", "Aluminio (Al)", "Martes", "Turno 1", "< 0,10 ppm"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "pH (Muestra 1)", "Todos", "Todos", "6,5 - 7,2"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "pH (Muestra 2)", "Todos", "Todos", "6,5 - 7,2"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Cloro residual", "Todos", "Turno 1", "< 0,10"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Turbiedad", "Todos", "Turno 1", "< 0,50"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Evaluación Sensorial", "Todos", "Todos", "Pasa / No Pasa"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Intensidad lampara UV 1", "Todos", "Todos", "> 25"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Color", "Miércoles", "Turno 1", "< 15 UPC"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Hierro total (Fe)", "Miércoles", "Turno 1", "< 0,10"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Dureza Total", "Miércoles", "Turno 1", "< 90 ppm"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Cloruros", "Miércoles", "Turno 1", "< 70 ppm"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Solidos disueltos totales", "Miércoles", "Turno 1", "< 500 ppm"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Conductividad", "Miércoles", "Turno 1", "< 745 us/cm"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Sulfatos", "Martes", "Turno 1", "< 50 ppm"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Manganeso", "Martes", "Turno 1", "< 0,050 ppm"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Alcalinidad Total", "Martes", "Turno 1", "< 90 ppm"),
    ("Tratamiento con luz Ultravioleta agua de servicio", "Aluminio (Al)", "Martes", "Turno 1", "< 0,10 ppm"),
    ("Cuarto generación dióxido de cloro", "Nivel de los tanques buffer", "Todos", "Todos", "%")
]

class ItemRegistro(BaseModel):
    zona: str
    parametro: str
    valor: str
    limite: str

class PayloadGuardarPlan(BaseModel):
    fecha: str
    turno: str
    operador: str
    registros: List[ItemRegistro]

class PayloadExcepcion(BaseModel):
    fecha: str
    turno: str
    zona: str
    parametro: str
    estado: str
    observacion: Optional[str] = ""

@app.get("/", response_class=HTMLResponse)
def vista_plan_calidad(request: Request):
    return templates.TemplateResponse(request, "plan_calidad.html", {
        "operadores": OPERADORES_PTAP,
        "maestra": MAESTRA_PLAN_CALIDAD
    })

@app.get("/api/plan-calidad/maestra")
def obtener_maestra(fecha: Optional[str] = None, turno: Optional[str] = None):
    """
    Filtra la maestra de plan de calidad según el día de la semana de la fecha 
    y el turno seleccionado, coincidiendo con la base de datos de Excel.
    """
    dia_semana_esp = "Todos"
    if fecha:
        try:
            dt = datetime.strptime(fecha, "%Y-%m-%d")
            dias_map = {
                0: "Lunes",
                1: "Martes",
                2: "Miércoles",
                3: "Jueves",
                4: "Viernes",
                5: "Sábado",
                6: "Domingo"
            }
            dia_semana_esp = dias_map.get(dt.weekday(), "Todos")
        except Exception:
            pass

    maestra_filtrada = []
    for zona, parametro, dia_esp, trn, limite in MAESTRA_PLAN_CALIDAD:
        coincide_dia = (dia_esp == "Todos" or dia_esp == dia_semana_esp)
        coincide_turno = True
        if turno and turno != "Todos":
            coincide_turno = (trn == "Todos" or trn == turno)

        if coincide_dia and coincide_turno:
            maestra_filtrada.append((zona, parametro, dia_esp, trn, limite))

    return {"maestra": maestra_filtrada}

@app.get("/api/plan-calidad/registros")
def consultar_registros(
    f_desde: str = Query(...),
    f_hasta: str = Query(...),
    turno: str = Query("Todos"),
    db: Session = Depends(get_db)
):
    query = db.query(models.RegistroPlanCalidad).filter(
        models.RegistroPlanCalidad.fecha >= f_desde,
        models.RegistroPlanCalidad.fecha <= f_hasta
    )
    if turno != "Todos":
        query = query.filter(models.RegistroPlanCalidad.turno == turno)
    
    registros = query.order_by(models.RegistroPlanCalidad.fecha.desc(), models.RegistroPlanCalidad.id.desc()).all()
    return {"total": len(registros), "data": registros}

@app.post("/api/plan-calidad/guardar")
def guardar_plan_calidad(payload: PayloadGuardarPlan, db: Session = Depends(get_db)):
    nuevos_modelos = []
    for item in payload.registros:
        val = item.valor.strip()
        if not val:
            continue
        
        resultado = "Cumple"
        if val.lower() in ["no pasa", "no cumple"]:
            resultado = "FUERA DE RANGO"
        
        reg = models.RegistroPlanCalidad(
            fecha=payload.fecha,
            turno=payload.turno,
            operador=payload.operador,
            zona=item.zona,
            parametro=item.parametro,
            valor=val,
            resultado=resultado,
            limite=item.limite
        )
        nuevos_modelos.append(reg)

    if not nuevos_modelos:
        raise HTTPException(status_code=400, detail="No se enviaron datos válidos para guardar.")

    db.add_all(nuevos_modelos)
    db.commit()

    sincronizar_excel(payload, nuevos_modelos)

    return {"status": "ok", "guardados": len(nuevos_modelos)}

@app.post("/api/plan-calidad/excepcion")
def guardar_excepcion(payload: PayloadExcepcion, db: Session = Depends(get_db)):
    exc = models.ExcepcionPlanCalidad(
        fecha=payload.fecha,
        turno=payload.turno,
        zona=payload.zona,
        parametro=payload.parametro,
        estado=payload.estado,
        observacion=payload.observacion
    )
    db.add(exc)
    db.commit()
    return {"status": "ok", "message": "Excepción guardada correctamente"}

def sincronizar_excel(payload: PayloadGuardarPlan, registros):
    nombre_archivo = "Bitacora_PTAP.xlsx"
    nombre_hoja = "Plan_Calidad"
    encabezados = ["Fecha", "Turno", "Operador", "Zona", "Parámetro", "Valor", "Resultado"]

    try:
        if os.path.exists(nombre_archivo):
            wb = openpyxl.load_workbook(nombre_archivo)
            if nombre_hoja in wb.sheetnames:
                ws = wb[nombre_hoja]
            else:
                ws = wb.create_sheet(title=nombre_hoja)
                ws.append(encabezados)
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = nombre_hoja
            ws.append(encabezados)

        for r in registros:
            ws.append([r.fecha, r.turno, r.operador, r.zona, r.parametro, r.valor, r.resultado])

        wb.save(nombre_archivo)
        wb.close()
    except Exception as e:
        print(f"Error al sincronizar con Excel: {e}")
